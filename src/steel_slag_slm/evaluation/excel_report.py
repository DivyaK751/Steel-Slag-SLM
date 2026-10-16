"""Write one evaluation run to an Excel workbook: Summary, Predictions, Disagreements."""

from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from steel_slag_slm.evaluation.metrics import FAILED, LABELS, EvalMetrics

BOLD = Font(bold=True)
WRAPPED_COLUMN_WIDTH = 60


@dataclass(frozen=True, slots=True)
class PredictionRow:
    """One paper in the report."""

    openalex_id: str
    title: str
    human_label: str | None
    agent_label: str
    match: str  # "✓", "✗", or "" when there is no human label
    relevance_score: float | None
    relevance_confidence: str | None
    relevance_reason: str
    supporting_excerpt: str | None


def _percent(value: float | None) -> str:
    return "n/a" if value is None else f"{value:.0%}"


def _write_table(sheet: Worksheet, header: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    sheet.append(list(header))
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD
    for row in rows:
        sheet.append(list(row))


def _write_summary(sheet: Worksheet, run_info: Mapping[str, object], metrics: EvalMetrics) -> None:
    for key, value in run_info.items():
        sheet.append([key, str(value)])
    sheet.append([])
    sheet.append(["Papers with a human label", metrics.labelled])
    sheet.append(["Agent agreed with human", metrics.matched])
    sheet.append(["Accuracy", _percent(metrics.accuracy)])
    sheet.append([])
    _write_table(
        sheet,
        [
            "Label",
            "Agent said it",
            "...and human agreed",
            "How often right",
            "Human said it",
            "...and agent caught it",
            "How many caught",
        ],
        [
            [
                s.label,
                s.agent_said,
                s.agent_correct,
                _percent(s.precision),
                s.human_said,
                s.agent_caught,
                _percent(s.recall),
            ]
            for s in metrics.per_label
        ],
    )
    sheet.append([])
    # Rows are the human's label, columns the agent's, so off-diagonal cells are mistakes
    agent_columns = [*LABELS, FAILED]
    _write_table(
        sheet,
        ["Human ↓ / Agent →", *agent_columns],
        [
            [human, *(metrics.confusion.get((human, agent), 0) for agent in agent_columns)]
            for human in LABELS
        ],
    )
    sheet.column_dimensions["A"].width = 32


def _write_predictions(sheet: Worksheet, rows: Sequence[PredictionRow]) -> None:
    header = list(PredictionRow.__dataclass_fields__)
    _write_table(sheet, header, [list(asdict(row).values()) for row in rows])
    for column in ("B", "H", "I"):  # title, reason, excerpt
        sheet.column_dimensions[column].width = WRAPPED_COLUMN_WIDTH
    sheet.freeze_panes = "A2"


def write_report(
    path: Path,
    run_info: Mapping[str, object],
    metrics: EvalMetrics,
    rows: Sequence[PredictionRow],
) -> Path:
    """Save the workbook to path (creating its folder) and return the path."""
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    _write_summary(summary, run_info, metrics)
    _write_predictions(workbook.create_sheet("Predictions"), rows)
    _write_predictions(
        workbook.create_sheet("Disagreements"), [row for row in rows if row.match == "✗"]
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
