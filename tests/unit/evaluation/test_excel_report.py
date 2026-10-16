from pathlib import Path

from openpyxl import load_workbook

from steel_slag_slm.evaluation.excel_report import PredictionRow, write_report
from steel_slag_slm.evaluation.metrics import compute_metrics


def _row(paper_id: str, human: str | None, agent: str, match: str) -> PredictionRow:
    return PredictionRow(paper_id, "t", human, agent, match, 0.9, "High", "r", None)


def test_write_report_has_three_sheets_and_only_mismatches_in_disagreements(
    tmp_path: Path,
) -> None:
    rows = [_row("W1", "Yes", "Yes", "✓"), _row("W2", "No", "Yes", "✗"), _row("W3", None, "No", "")]
    path = write_report(
        tmp_path / "out" / "run.xlsx",
        {"Model": "m"},
        compute_metrics([("Yes", "Yes"), ("No", "Yes")]),
        rows,
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Predictions", "Disagreements"]
    assert workbook["Predictions"].max_row == 4  # header + 3 papers
    disagreements = list(workbook["Disagreements"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in disagreements] == ["W2"]
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True) if row}
    assert summary["Accuracy"] == "50%"
