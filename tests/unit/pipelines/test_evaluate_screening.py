import asyncio
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

import steel_slag_slm.pipelines.evaluate_screening as evaluate_screening
from steel_slag_slm.config import EvaluationSettings, ScreeningSettings
from steel_slag_slm.exceptions import AgentOutputError
from steel_slag_slm.llm.prompts import load_prompt
from steel_slag_slm.schemas import PaperForScreening, RelevanceDecision

LABELLED_ROWS = [
    {
        "openalex_id": "W1",
        "title": "a",
        "abstract": "x",
        "year": 2020,
        "keywords": "slag, cement",
        "human_label": "Yes",
    },
    {
        "openalex_id": "W2",
        "title": "b",
        "abstract": "x",
        "year": 2020,
        "keywords": None,
        "human_label": "no",
    },
    {
        "openalex_id": "W3",
        "title": "c",
        "abstract": "x",
        "year": 2020,
        "keywords": None,
        "human_label": None,
    },
    {
        "openalex_id": "W4",
        "title": "d",
        "abstract": "x",
        "year": 2020,
        "keywords": None,
        "human_label": "Unclear",
    },
]


def _run(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    seen_keywords: dict[str, tuple[str, ...]] = {}

    async def fake_screen(_runner: Any, _prompt: Any, paper: PaperForScreening) -> Any:
        seen_keywords[paper.openalex_id] = paper.keywords
        if paper.openalex_id == "W4":
            raise AgentOutputError("bad json")
        return RelevanceDecision(
            relevant="Yes", relevance_score=0.9, relevance_confidence="High", relevance_reason="r"
        )

    monkeypatch.setattr(evaluate_screening, "query_rows", lambda *_: LABELLED_ROWS)
    monkeypatch.setattr(evaluate_screening, "screen_paper", fake_screen)
    screening = ScreeningSettings(
        model="m",
        temperature=0.0,
        prompt_name="screen_relevance",
        batch_size=10,
        max_concurrency=2,
        papers_table="p.d.papers",
    )
    evaluation = EvaluationSettings(labels_table="p.d.labels", output_dir=tmp_path)
    path, metrics = asyncio.run(
        evaluate_screening.run_evaluation(
            screening,
            evaluation,
            None,
            None,
            load_prompt("screen_relevance"),
            {"Model": "m"},  # type: ignore[arg-type]
        )
    )
    return path, metrics, seen_keywords


def test_run_evaluation_scores_only_labelled_papers_and_keeps_failures(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path, metrics, seen_keywords = _run(monkeypatch, tmp_path)
    # W3 has no label; W4 failed and counts as a miss
    assert (metrics.labelled, metrics.matched) == (3, 1)
    assert seen_keywords["W1"] == ("slag", "cement")
    predictions = list(load_workbook(path)["Predictions"].iter_rows(min_row=2, values_only=True))
    assert [(row[0], row[3], row[4]) for row in predictions] == [
        ("W1", "Yes", "✓"),
        ("W2", "Yes", "✗"),
        ("W3", "Yes", None),  # Excel stores an empty cell as None
        ("W4", "Failed", "✗"),
    ]
    assert path.name.startswith("eval_") and path.name.endswith("_m_prompt-v1.xlsx")
