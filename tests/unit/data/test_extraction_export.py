from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

from steel_slag_slm.data.extraction_export import export_workbook
from steel_slag_slm.data.extraction_tables import SHEET_TABLES


def test_export_workbook_writes_one_sorted_sheet_per_table(tmp_path: Path) -> None:
    client = MagicMock()
    client.get_table.side_effect = lambda table_id: SimpleNamespace(
        table_id=table_id,
        schema=[SimpleNamespace(name="source_id"), SimpleNamespace(name="extracted_at")],
    )
    stamp = datetime(2026, 10, 25, 12, 0, tzinfo=UTC)
    rows = [{"source_id": "W2", "extracted_at": stamp}, {"source_id": "W1", "extracted_at": stamp}]
    client.list_rows.side_effect = lambda table: (
        rows if table.table_id.endswith("sheet_01_scientific_literature") else []
    )
    output = tmp_path / "out" / "extraction.xlsx"

    counts = export_workbook(client, "p.d", output)

    workbook = load_workbook(output)
    assert workbook.sheetnames == [sheet_table.table for sheet_table in SHEET_TABLES]
    literature = workbook["sheet_01_scientific_literature"]
    assert [cell.value for cell in literature[1]] == ["source_id", "extracted_at"]
    assert [literature.cell(row, 1).value for row in (2, 3)] == ["W1", "W2"]
    assert literature.cell(2, 2).value == datetime(2026, 10, 25, 12, 0)
    assert counts["sheet_01_scientific_literature"] == 2
    assert counts["sheet_09_citations_provenance"] == 0
